
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parent
INPUT_DIR = PROJECT_DIR / "data" / "input"
OUTPUT_DIR = PROJECT_DIR / "output"
REPORT_PATH = OUTPUT_DIR / "sales_report.xlsx"

REQUIRED_COLUMNS = [
    "order_id",
    "order_date",
    "product",
    "category",
    "region",
    "quantity",
    "unit_price",
]


def load_sales_files(input_dir):
    """Read and combine every Excel file in the input folder."""
    files = sorted(Path(input_dir).glob("*.xlsx"))

    if not files:
        raise FileNotFoundError(
            f"No .xlsx files found in {input_dir}"
        )

    frames = []

    for file_path in files:
        df = pd.read_excel(file_path)
        df.columns = (
            df.columns.astype(str)
            .str.strip()
            .str.lower()
            .str.replace(" ", "_")
        )

        missing = set(REQUIRED_COLUMNS) - set(df.columns)

        if missing:
            raise ValueError(
                f"{file_path.name} is missing columns: "
                f"{', '.join(sorted(missing))}"
            )

        df = df[REQUIRED_COLUMNS].copy()
        frames.append(df)

    return pd.concat(frames, ignore_index=True)


def prepare_sales_data(raw_df):
    """Remove duplicates and separate valid and invalid records."""
    df = raw_df.copy()

    df["order_date"] = pd.to_datetime(
        df["order_date"], errors="coerce"
    )

    for column in ["order_id", "product", "category", "region"]:
        df[column] = df[column].astype("string").str.strip()
        df[column] = df[column].replace("", pd.NA)

    df["quantity"] = pd.to_numeric(
        df["quantity"], errors="coerce"
    )
    df["unit_price"] = pd.to_numeric(
        df["unit_price"], errors="coerce"
    )

    duplicate_mask = df.duplicated(
        subset=REQUIRED_COLUMNS, keep="first"
    )
    duplicate_count = int(duplicate_mask.sum())
    df = df.loc[~duplicate_mask].copy()

    issues = pd.Series("", index=df.index, dtype="string")

    required_fields = [
        "order_id",
        "order_date",
        "product",
        "category",
        "region",
        "quantity",
        "unit_price",
    ]

    for column in required_fields:
        missing_mask = df[column].isna()
        issues.loc[missing_mask] += f"Missing {column}; "

    bad_quantity = df["quantity"].notna() & (df["quantity"] <= 0)
    issues.loc[bad_quantity] += "Quantity must be greater than zero; "

    bad_price = df["unit_price"].notna() & (df["unit_price"] < 0)
    issues.loc[bad_price] += "Unit price cannot be negative; "

    invalid_mask = issues.ne("")

    invalid_df = df.loc[invalid_mask].copy()
    invalid_df["validation_issues"] = issues.loc[invalid_mask]

    valid_df = df.loc[~invalid_mask].copy()
    valid_df["revenue"] = (
        valid_df["quantity"] * valid_df["unit_price"]
    ).round(2)

    return valid_df, invalid_df, duplicate_count


def create_report_tables(valid_df, invalid_df, duplicate_count):
    """Create summary metrics and sales breakdowns."""
    monthly = (
        valid_df.assign(
            month=valid_df["order_date"].dt.to_period("M").astype(str)
        )
        .groupby("month", as_index=False)["revenue"]
        .sum()
        .round(2)
    )

    products = (
        valid_df.groupby("product", as_index=False)
        .agg(
            units_sold=("quantity", "sum"),
            revenue=("revenue", "sum"),
        )
        .sort_values("revenue", ascending=False)
        .reset_index(drop=True)
    )
    products["revenue"] = products["revenue"].round(2)

    regions = (
        valid_df.groupby("region", as_index=False)
        .agg(
            orders=("order_id", "nunique"),
            revenue=("revenue", "sum"),
        )
        .sort_values("revenue", ascending=False)
        .reset_index(drop=True)
    )
    regions["revenue"] = regions["revenue"].round(2)

    summary = pd.DataFrame(
        {
            "Metric": [
                "Input records",
                "Duplicate records removed",
                "Valid sales records",
                "Records requiring review",
                "Unique orders",
                "Total revenue",
            ],
            "Value": [
                len(valid_df) + len(invalid_df) + duplicate_count,
                duplicate_count,
                len(valid_df),
                len(invalid_df),
                valid_df["order_id"].nunique(),
                round(valid_df["revenue"].sum(), 2),
            ],
        }
    )

    return summary, monthly, products, regions


def format_workbook(writer):
    """Apply simple professional formatting and add charts."""
    workbook = writer.book

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid", fgColor="1F4E78"
            )

        for column_cells in worksheet.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            column_letter = get_column_letter(column_cells[0].column)
            worksheet.column_dimensions[column_letter].width = min(
                max_length + 3, 35
            )

    monthly_sheet = workbook["Monthly Sales"]
    if monthly_sheet.max_row > 1:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Revenue by Month"
        chart.y_axis.title = "Revenue"
        chart.x_axis.title = "Month"

        chart.add_data(
            Reference(
                monthly_sheet,
                min_col=2,
                min_row=1,
                max_row=monthly_sheet.max_row,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                monthly_sheet,
                min_col=1,
                min_row=2,
                max_row=monthly_sheet.max_row,
            )
        )
        monthly_sheet.add_chart(chart, "D2")

    product_sheet = workbook["Product Sales"]
    if product_sheet.max_row > 1:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Revenue by Product"
        chart.x_axis.title = "Revenue"
        chart.y_axis.title = "Product"

        chart.add_data(
            Reference(
                product_sheet,
                min_col=3,
                min_row=1,
                max_row=product_sheet.max_row,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                product_sheet,
                min_col=1,
                min_row=2,
                max_row=product_sheet.max_row,
            )
        )
        product_sheet.add_chart(chart, "E2")


def generate_report(valid_df, invalid_df, duplicate_count, report_path):
    """Write all report sheets to one Excel workbook."""
    summary, monthly, products, regions = create_report_tables(
        valid_df, invalid_df, duplicate_count
    )

    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        monthly.to_excel(writer, sheet_name="Monthly Sales", index=False)
        products.to_excel(writer, sheet_name="Product Sales", index=False)
        regions.to_excel(writer, sheet_name="Regional Sales", index=False)
        valid_df.to_excel(writer, sheet_name="Cleaned Sales", index=False)
        invalid_df.to_excel(writer, sheet_name="Invalid Records", index=False)

        format_workbook(writer)

    return summary


def main():
    try:
        raw_df = load_sales_files(INPUT_DIR)

        valid_df, invalid_df, duplicate_count = prepare_sales_data(
            raw_df
        )

        summary = generate_report(
            valid_df,
            invalid_df,
            duplicate_count,
            REPORT_PATH,
        )

        print("\nSales report generated successfully!")
        print(f"Report location: {REPORT_PATH}")
        print("\nSummary:")
        print(summary.to_string(index=False))

    except (FileNotFoundError, ValueError) as error:
        print(f"Error: {error}")


if __name__ == "__main__":
    main()
